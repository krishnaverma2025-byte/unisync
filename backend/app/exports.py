from io import BytesIO
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from reportlab.platypus import SimpleDocTemplate,Paragraph,Table,TableStyle,Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.pagesizes import A4,landscape
from .slots import DAYS

def time(n): return f'{n//60:02d}:{n%60:02d}'
HEADERS=['Day','Start','End','Subject','Codes','Section','Faculty','Room']
def values(entries):
    return [[DAYS[e['day']],time(e['start']),time(e['end']),e['subject'],e['codes'],e['section'],e['faculty'],e['room']] for e in sorted(entries,key=lambda e:(e['day'],e['start'],e['section']))]
def export_xlsx(title,entries):
    book=Workbook();sheet=book.active;sheet.title='Timetable';sheet.append(["'"+title if title[:1] in ('=','+','-','@') else title]);sheet.append(HEADERS)
    for row in values(entries):
        # Untrusted names are always strings; prevent Excel formula injection.
        sheet.append(["'"+v if isinstance(v,str) and v[:1] in ('=','+','-','@') else v for v in row])
    for cell in sheet[2]: cell.fill=PatternFill('solid',fgColor='295B43');cell.font=Font(color='FFFFFF',bold=True)
    widths=[15,10,10,30,20,18,25,25]
    for index,width in enumerate(widths,1): sheet.column_dimensions[sheet.cell(2,index).column_letter].width=width
    sheet.freeze_panes='A3';sheet.auto_filter.ref=f'A2:H{sheet.max_row}'
    for row in sheet.iter_rows(min_row=3):
        for cell in row: cell.alignment=Alignment(vertical='top',wrap_text=True)
    output=BytesIO();book.save(output);return output.getvalue()
def export_pdf(title,entries):
    output=BytesIO();doc=SimpleDocTemplate(output,pagesize=landscape(A4),leftMargin=25,rightMargin=25,topMargin=25,bottomMargin=25)
    styles=getSampleStyleSheet();styles['BodyText'].fontSize=8;styles['BodyText'].leading=11
    data=[[Paragraph(escape(v),styles['BodyText']) for v in row] for row in [HEADERS]+values(entries)]
    table=Table(data,colWidths=[60,40,40,140,75,75,120,120],repeatRows=1)
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#dfead3')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f4f7ee')]),('VALIGN',(0,0),(-1,-1),'TOP'),('LINEBELOW',(0,0),(-1,0),.8,colors.HexColor('#295b43')),('LEFTPADDING',(0,0),(-1,-1),6),('RIGHTPADDING',(0,0),(-1,-1),6),('TOPPADDING',(0,0),(-1,-1),8),('BOTTOMPADDING',(0,0),(-1,-1),8)]))
    doc.build([Paragraph(escape(title),styles['Title']),Paragraph('Recurring Monday–Friday · exact campus clock intervals · Theory and Lab are distinct systems',styles['BodyText']),Spacer(1,15),table]);return output.getvalue()
