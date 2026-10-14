from io import BytesIO
import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine,text,inspect
from conftest import headers
from app.models import *
from app.engine import candidates,solve,SchedulingFailure,validate_entries,generate
from app.analytics import evaluate
from app.migrations import migrate
from app.natural_language import parse

def isolate(db):
    db.query(Assignment).filter(Assignment.id!=1).delete()
    a=db.get(Assignment,1);a.weekly_sessions=1
    db.query(SchedulingConstraint).delete();db.commit();return a

def test_unavailable_overrides_preferred_and_working_days(db):
    a=isolate(db)
    db.add(Availability(faculty_id=a.faculty_id,day=0,start=480,end=600,status='preferred'))
    db.add(Availability(faculty_id=a.faculty_id,day=0,start=500,end=550,status='unavailable'))
    db.query(WorkingDay).filter_by(day=1).update({'enabled':False});db.commit()
    options=candidates(db,a)
    assert not any(o['day']==0 and o['start']<550 and o['end']>500 for o in options)
    assert not any(o['day']==1 for o in options)

def test_preferences_and_hard_time_rules_change_solution(db):
    a=isolate(db)
    db.add(Availability(faculty_id=a.faculty_id,day=3,start=840,end=890,status='preferred'))
    db.add(SchedulingConstraint(name='Faculty preferences',rule='faculty_preferences',severity='SOFT',value=0,weight=100))
    db.add(SchedulingConstraint(name='Afternoon only',rule='pm_only',severity='HARD',faculty_id=a.faculty_id,value=0,weight=100))
    db.commit();solutions,_=solve(db,time_limit=2)
    entry=solutions[0]['entries'][0]
    assert entry['day']==3 and entry['start']==840
    assert evaluate(db,solutions[0]['entries'])['faculty_satisfaction']==100

def test_alternatives_are_distinct_valid_and_scored(db):
    isolate(db)
    solutions,_=solve(db,count=3,time_limit=2)
    assert len(solutions)==3
    signatures={tuple((e['assignment_id'],e['day'],e['start']) for e in s['entries']) for s in solutions}
    assert len(signatures)==3
    for s in solutions:
        validate_entries(db,s['entries'],complete=True)
        score=evaluate(db,s['entries'])
        assert score['conflict_count']==0
        assert 0<=score['quality_score']<=100
        assert sum(c['weight'] for c in score['components'].values())==100

def test_weekly_limits_are_explained(db):
    f=db.get(Faculty,1);f.max_classes_week=1;db.commit()
    with pytest.raises(SchedulingFailure) as error: solve(db,time_limit=1)
    assert any('limits' in i['message'] for i in error.value.detail['issues'])

def test_infeasibility_core_from_clock_collisions(db):
    db.query(SchedulingConstraint).delete();db.query(Availability).delete()
    db.query(Assignment).filter(Assignment.id.notin_([1,2])).delete()
    for a in db.query(Assignment):
        a.weekly_sessions=1
        db.add(Availability(faculty_id=a.faculty_id,day=0,start=480,end=530,status='available'))
    db.commit()
    with pytest.raises(SchedulingFailure) as error: solve(db,time_limit=2)
    assert error.value.detail['status']=='INFEASIBLE'
    assert error.value.detail['issues']
    assert any('CSE-A' in i['message'] for i in error.value.detail['issues'])

def test_local_natural_language_preview_is_explicit_and_safe(db):
    result=parse(db,'Dr. Meera should preferably have no classes after 3 PM and Programming Lab should happen in the morning')
    assert len(result['constraints'])==2 and not result['warnings']
    assert result['constraints'][0]['rule']=='end_before' and result['constraints'][0]['value']==900
    assert result['constraints'][0]['severity']=='SOFT'
    assert result['constraints'][1]['rule']=='am_only' and result['constraints'][1]['severity']=='HARD'
    assert parse(db,'Dr. Unknown should have no classes after 3 PM')['warnings']
    assert not parse(db,'Dr. Unknown should have no classes after 3 PM')['constraints']
    assert parse(db,'everyone should have no classes after 25:80')['warnings']

def test_additive_migration_preserves_existing_rows(tmp_path):
    engine=create_engine('sqlite:///'+str(tmp_path/'legacy.db'))
    with engine.begin() as c:
        c.execute(text('CREATE TABLE faculty (id INTEGER PRIMARY KEY, name VARCHAR, department_id INTEGER, user_id INTEGER)'))
        c.execute(text("INSERT INTO faculty VALUES (99,'Existing Faculty',NULL,NULL)"))
    assert migrate(engine) is True
    assert migrate(engine) is False
    with engine.connect() as c:
        row=c.execute(text('SELECT name,max_classes_day,max_classes_week FROM faculty WHERE id=99')).one()
        assert tuple(row)==('Existing Faculty',6,24)
    assert 'constraints' in inspect(engine).get_table_names()

def test_student_section_exports_and_free_rooms(client,db):
    h=headers(client);s=headers(client,'krishna@unisync.local')
    isolate(db)
    table=generate(db,'Export demo');table.published=True;db.commit()
    assert client.put('/student/section',headers=s,json={'section_id':1}).status_code==200
    assert client.get('/auth/me',headers=s).json()['section_id']==1
    stats=client.get(f'/timetables/{table.id}/analytics',headers=s).json()
    assert stats['faculty'] and stats['rooms'] and stats['conflict_count']==0
    pdf=client.get(f'/timetables/{table.id}/export/pdf?view=section&entity_id=1',headers=s)
    assert pdf.status_code==200 and pdf.content.startswith(b'%PDF')
    xlsx=client.get(f'/timetables/{table.id}/export/xlsx',headers=s)
    assert xlsx.status_code==200
    book=load_workbook(BytesIO(xlsx.content));assert book.active.max_row==3
    e=db.query(Entry).first()
    r=client.post('/free-rooms',headers=s,json={'day':e.day,'start':e.start,'end':e.end,'capacity':1}).json()
    assert e.room_id not in {room['id'] for room in r['rooms']}

def test_validate_and_move_are_atomic_with_conflict_messages(client,db):
    h=headers(client)
    db.query(SchedulingConstraint).delete();db.query(Assignment).filter(Assignment.id.notin_([1,2])).delete()
    for a in db.query(Assignment): a.weekly_sessions=1
    db.commit();table=generate(db,'Editable')
    entries=db.query(Entry).filter_by(timetable_id=table.id).all();a,b=entries
    before=(a.day,a.start,a.codes)
    payload={'entry_id':a.id,'day':b.day,'start':b.start,'room_id':a.room_id}
    response=client.post('/admin/validate-change',headers=h,json=payload)
    assert response.status_code==409 and 'clash' in response.json()['detail']
    assert (db.get(Entry,a.id).day,db.get(Entry,a.id).start,db.get(Entry,a.id).codes)==before
    candidate=next(o for o in candidates(db,db.get(Assignment,a.assignment_id)) if o['day']!=b.day and o['room_id']==a.room_id)
    payload.update(day=candidate['day'],start=candidate['start'])
    assert client.put(f'/admin/entries/{a.id}/move',headers=h,json=payload).status_code==200
    assert db.get(Entry,a.id).start==candidate['start'] and db.get(Entry,a.id).day==candidate['day']

def test_faculty_request_review_and_ownership(client,db):
    admin_h=headers(client);faculty_h=headers(client,'meera@unisync.local')
    isolate(db);table=generate(db,'Faculty');table.published=True;db.commit();e=db.query(Entry).one()
    workload=client.get('/faculty/workload',headers=faculty_h)
    assert workload.status_code==200 and workload.json()['workload']['periods']==1
    response=client.post('/change-requests',headers=faculty_h,json={'timetable_id':table.id,'entry_id':e.id,'reason':'Please move this class later.'})
    assert response.status_code==200;rid=response.json()['id']
    assert client.patch(f'/admin/change-requests/{rid}',headers=faculty_h,json={'status':'approved'}).status_code==403
    assert client.patch(f'/admin/change-requests/{rid}',headers=admin_h,json={'status':'approved','admin_note':'Will review the draft.'}).status_code==200
    assert client.get('/change-requests',headers=faculty_h).json()[0]['admin_note']=='Will review the draft.'

def test_invalid_rules_and_roles_rejected(client):
    h=headers(client);student=headers(client,'krishna@unisync.local')
    assert client.post('/admin/constraints/parse',headers=student,json={'text':'everyone in the morning'}).status_code==403
    assert client.post('/admin/constraints',headers=h,json={'name':'Bad','rule':'make_everything_perfect','severity':'SOFT','value':1}).status_code==422
    assert client.post('/admin/constraints',headers=h,json={'name':'Bad','rule':'student_gaps','severity':'HARD'}).status_code==422

def test_alternative_api_ranks_and_preserves_publication(client,db):
    isolate(db);h=headers(client)
    response=client.post('/admin/alternatives',headers=h,json={'name':'Comparison','count':3,'time_limit':2})
    assert response.status_code==200
    data=response.json();scores=[s['quality']['quality_score'] for s in data['alternatives']]
    assert len(scores)==3 and scores==sorted(scores,reverse=True)
    assert data['best_id']==data['alternatives'][0]['id']
    assert db.query(Timetable).filter_by(published=True).count()==0
    assert db.query(Entry).count()==3

def test_batch_rules_atomic_and_published_guard(client,db):
    isolate(db);h=headers(client);table=generate(db,'Protected');table.published=True;db.commit()
    original=db.query(SchedulingConstraint).count()
    response=client.post('/admin/constraints/batch',headers=h,json={'constraints':[
        {'name':'Soft late','rule':'avoid_last','severity':'SOFT'},
        {'name':'Invalid rule','rule':'unsupported','severity':'SOFT'}]})
    assert response.status_code==422 and db.query(SchedulingConstraint).count()==original
    e=db.query(Entry).one()
    response=client.post('/admin/constraints/batch',headers=h,json={'constraints':[{'name':'Exclude saved class','rule':'start_after','severity':'HARD','value':e.end+1,'faculty_id':1}]})
    assert response.status_code==409 and db.query(SchedulingConstraint).count()==original
    assert db.query(Entry).one().start==e.start

def test_faculty_preferences_edit_preserves_existing_published_entries(client,db):
    isolate(db);table=generate(db,'Stable');table.published=True;db.commit()
    h=headers(client,'meera@unisync.local')
    response=client.put('/faculty/availability',headers=h,json={'windows':[{'day':day,'start':480,'end':1160,'status':'preferred'} for day in range(5)]})
    assert response.status_code==200
    assert db.query(Entry).count()==1 and db.query(Timetable).one().published
    e=db.query(Entry).one()
    response=client.put('/faculty/availability',headers=h,json={'windows':[{'day':e.day,'start':e.start,'end':e.end,'status':'unavailable'}]})
    assert response.status_code==409
    assert db.query(Availability).filter_by(faculty_id=1,status='preferred').count()==5

def test_hard_consecutive_limit_and_daily_cap(db):
    a=isolate(db);a.weekly_sessions=2
    db.query(Availability).delete();db.add(Availability(faculty_id=1,day=0,start=480,end=590,status='available'))
    db.add(SchedulingConstraint(name='No back-to-back lectures',rule='max_consecutive',severity='HARD',faculty_id=1,value=1,weight=100))
    db.commit()
    with pytest.raises(SchedulingFailure): solve(db,time_limit=2)
    db.query(SchedulingConstraint).delete();db.get(Faculty,1).max_classes_day=1;db.commit()
    with pytest.raises(SchedulingFailure): solve(db,time_limit=2)
